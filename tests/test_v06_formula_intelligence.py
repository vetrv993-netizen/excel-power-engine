from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from excel_power_engine.formula_intelligence import cell_info, trace_formula, formula_issues, extract_references


def make_book(path: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "تحليل"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["B1"] = "=A1+A2"
    ws["B2"] = "='تحليل'!B1+5"
    ws["B3"] = "=A1+#REF!"
    wb.save(path)


def test_cell_info_and_formula_trace(tmp_path):
    p = tmp_path / "demo.xlsx"
    make_book(p)
    snap = cell_info(p, "تحليل", "B1")
    assert snap.cell_type == "formula"
    assert snap.formula == "A1+A2"
    assert snap.cached_value is None or snap.cached_value == 30
    trace = trace_formula(p, "تحليل", "B1")
    assert "تحليل!A1" in [x["cell"] for x in trace["dependencies"]]
    assert "تحليل!A2" in [x["cell"] for x in trace["dependencies"]]


def test_formula_issue_scan_and_reference_parser(tmp_path):
    p = tmp_path / "demo.xlsx"
    make_book(p)
    refs = extract_references("='تحليل'!B1+A1", "تحليل")
    assert "تحليل!B1" in refs
    assert "تحليل!A1" in refs
    issues = formula_issues(p)
    assert any(x["cell"] == "تحليل!B3" and x["kind"] == "broken-reference" for x in issues)
