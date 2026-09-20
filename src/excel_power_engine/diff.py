from __future__ import annotations

from pathlib import Path
from typing import Any


def compare_cells(before: str | Path, after: str | Path, sheet: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb1 = load_workbook(before, data_only=False, keep_vba=True)
    wb2 = load_workbook(after, data_only=False, keep_vba=True)
    ws1 = wb1[sheet]
    ws2 = wb2[sheet]

    max_row = max(ws1.max_row, ws2.max_row)
    max_col = max(ws1.max_column, ws2.max_column)
    result = []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v1 = ws1.cell(r, c).value
            v2 = ws2.cell(r, c).value
            if v1 != v2:
                result.append({
                    "cell": ws1.cell(r, c).coordinate,
                    "before": v1,
                    "after": v2,
                })
    return result
