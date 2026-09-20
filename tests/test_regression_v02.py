from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from excel_power_engine.inspect import inspect_workbook
from excel_power_engine.safe_edit import EditOperation, SafeEditor


def make_sample(path):
    import openpyxl
    from openpyxl.worksheet.table import Table,TableStyleInfo
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Data';ws['A1']='Name';ws['B1']='Value';ws['A2']='Alpha';ws['B2']=10;ws['A3']='Beta';ws['B3']=20
    t=Table(displayName='DataTable',ref='A1:B3');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showFirstColumn=False,showLastColumn=False);ws.add_table(t);wb.save(path)


def inject(path):
    tmp=path.with_name('tmp.xlsx')
    with ZipFile(path,'r') as zin,ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            d=zin.read(i.filename)
            if i.filename=='xl/worksheets/sheet1.xml':
                txt=d.decode();ext='<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations count="1"><x14:dataValidation type="list"><x14:formula1><xm:f>$J$2:$J$14</xm:f></x14:formula1><xm:sqref>B2:B3</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>';d=txt.replace('</worksheet>',ext+'</worksheet>').encode();zout.writestr(i,d)
            else:zout.writestr(i,d)
        zout.writestr('xl/vbaProject.bin',b'FAKE-VBA-PAYLOAD')
    tmp.replace(path)


def test_deep_and_safe(tmp_path):
    p=tmp_path/'sample.xlsx';make_sample(p);inject(p);prof=inspect_workbook(p,deep=True);assert prof.has_vba and 'tables' in prof.features and prof.sheet_profiles['Data']['x14_data_validations']==1
    out=tmp_path/'edited.xlsx';r=SafeEditor().edit_cells(p,'Data',[EditOperation('B2',99),EditOperation('A3',value='Changed')],out,True);assert out.exists();assert r.vba_preserved and r.x14_preserved and r.protected_parts_unchanged;assert set(r.changed_parts)=={'xl/worksheets/sheet1.xml'}
    with ZipFile(out) as z: assert 'x14:dataValidation' in z.read('xl/worksheets/sheet1.xml').decode() and z.read('xl/vbaProject.bin')==b'FAKE-VBA-PAYLOAD'
