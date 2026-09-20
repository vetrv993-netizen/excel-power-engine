from __future__ import annotations

import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable
import warnings

from openpyxl import load_workbook


@dataclass
class SearchHit:
    sheet: str
    cell: str
    value: object
    formula: str | None = None
    matched_in: str = "value"

    def as_dict(self) -> dict:
        return asdict(self)


def _matcher(term: str, mode: str, case_sensitive: bool):
    flags = 0 if case_sensitive else re.IGNORECASE
    needle = term if case_sensitive else term.casefold()
    if mode == "regex":
        return lambda s: re.search(term, s, flags) is not None
    if mode == "exact":
        return lambda s: s == term if case_sensitive else s.casefold() == needle
    if mode == "starts":
        return lambda s: s.startswith(term) if case_sensitive else s.casefold().startswith(needle)
    if mode == "ends":
        return lambda s: s.endswith(term) if case_sensitive else s.casefold().endswith(needle)
    return lambda s: needle in (s if case_sensitive else s.casefold())


def search_workbook(
    path: str | Path,
    term: str,
    *,
    sheets: Iterable[str] | None = None,
    mode: str = "contains",
    case_sensitive: bool = False,
    search_values: bool = True,
    search_formulas: bool = True,
    max_results: int = 2000,
) -> list[SearchHit]:
    if mode not in {"contains", "exact", "starts", "ends", "regex"}:
        raise ValueError(f"Unsupported search mode: {mode}")
    matcher = _matcher(term, mode, case_sensitive)
    target_sheets = set(sheets) if sheets else None
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    hits: list[SearchHit] = []
    try:
        for ws in wb.worksheets:
            if target_sheets and ws.title not in target_sheets:
                continue
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    formula = value if isinstance(value, str) and value.startswith("=") else None
                    if search_formulas and formula is not None and matcher(formula):
                        hits.append(SearchHit(ws.title, cell.coordinate, value, formula, "formula"))
                    elif search_values:
                        text = str(value)
                        if matcher(text):
                            hits.append(SearchHit(ws.title, cell.coordinate, value, formula, "value"))
                    if len(hits) >= max_results:
                        return hits
        return hits
    finally:
        try:
            wb.close()
        except Exception:
            pass
