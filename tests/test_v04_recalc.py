from pathlib import Path

from excel_power_engine.recalculate import smart_recalculate
from excel_power_engine.smart_engine import build_plan


def make_simple(path: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=SUM(A1:A2)"
    wb.save(path)


def test_recalc_plan_without_native(monkeypatch, tmp_path):
    p = tmp_path / "sample.xlsx"
    make_simple(p)
    monkeypatch.setattr("excel_power_engine.smart_engine._native_excel_available", lambda: False)
    plan = build_plan(p, "recalculate", recalculate=True)
    assert plan.engine == "openpyxl"
    assert plan.requested_recalculate is True
    assert plan.warnings


def test_smart_recalc_does_not_claim_without_excel(monkeypatch, tmp_path):
    p = tmp_path / "sample.xlsx"
    make_simple(p)
    monkeypatch.setattr("excel_power_engine.smart_engine._native_excel_available", lambda: False)
    result = smart_recalculate(p)
    assert result["executed_engine"] == "openpyxl"
    assert result["errors"]
    assert "إعادة الحساب الفعلية" in result["errors"][0]


def test_native_plan_when_available(monkeypatch, tmp_path):
    p = tmp_path / "sample.xlsx"
    make_simple(p)
    monkeypatch.setattr("excel_power_engine.smart_engine._native_excel_available", lambda: True)
    plan = build_plan(p, "recalculate", recalculate=True)
    assert plan.engine == "excel-native"
    assert plan.requested_recalculate is True
