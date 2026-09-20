from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook

from excel_power_engine.target_engine import parse_target_specs
from excel_power_engine.search_engine import search_workbook
from excel_power_engine.workflow import validate_workflow


def make_book(tmp_path: Path) -> Path:
    p=tmp_path/'sample.xlsx'
    wb=Workbook(); ws=wb.active; ws.title='Data'
    ws['A1']='Ahmed'; ws['B1']='التهاب الضرع'; ws['C1']='=1+2'
    ws['A2']='Muhammad'; ws['B2']='ضرع'; ws['C2']='=C1+4'
    wb.save(p); return p


def test_target_parser():
    specs=parse_target_specs('A1:C3;11:12;F:H;V266')
    assert [s.kind for s in specs]==['range','row','column','cell']


def test_search_contains_and_formula(tmp_path):
    p=make_book(tmp_path)
    hits=search_workbook(p,'ضرع')
    assert {(h.sheet,h.cell) for h in hits}=={('Data','B1'),('Data','B2')}
    fh=search_workbook(p,'C1',search_values=False,search_formulas=True)
    assert any(h.cell=='C2' for h in fh)


def test_workflow_validation():
    ops=[{'kind':'format','sheet':'Data','ranges':['A1:C2'],'payload':{'fill':'#FFF2CC'}}]
    assert validate_workflow(ops)==ops
