from __future__ import annotations
import argparse, json
from .engine import ExcelEngine
from .safe_edit import EditOperation
from .recalculate import smart_recalculate
from .search_engine import search_workbook
from .target_engine import parse_target_specs, resolve_target_specs
from .format_engine import ExcelOperations
from .workflow import load_workflow, validate_workflow, run_workflow
from .operations import preview_operations
from .orchestrator import preview_workspace, run_workspace, rollback_workspace


def emit(x):
    print(json.dumps(x, ensure_ascii=False, indent=2, default=str))


def parse_edits(items):
    ops=[]
    for s in items:
        parts=s.split('::',2)
        if len(parts)==2:
            ops.append(EditOperation(parts[0], value=parts[1]))
        elif len(parts)==3 and parts[1].upper()=='FORMULA':
            ops.append(EditOperation(parts[0], formula=parts[2]))
        else:
            raise SystemExit(f'Invalid edit: {s}')
    return ops


def cmd_inspect(path, deep):
    p=ExcelEngine().profile(path, deep=deep)
    d={'path':str(p.path),'suffix':p.suffix,'ooxml':p.is_ooxml,'xlsm':p.is_xlsm,'has_vba':p.has_vba,'features':sorted(p.features),'sheets':p.sheet_parts,'parts':len(p.internal_parts)}
    if deep:d.update({'part_sizes':p.part_sizes,'defined_names':p.defined_names,'sheet_profiles':p.sheet_profiles,'integrity':p.integrity})
    emit(d)


def cmd_plan(path, operation, recalculate=False, prefer_native=False):
    emit(ExcelEngine().plan(path, operation, recalculate=recalculate, prefer_native=prefer_native).as_dict())


def cmd_read(path, sheet, engine):
    from .reader import read_dataframe
    df=read_dataframe(path, sheet=sheet, engine=engine)
    print(df.head(20).to_string(index=False))
    print('\nshape:',df.shape)


def cmd_diff(before, after, sheet):
    from openpyxl import load_workbook
    from .diff import compare_cells
    if sheet is None: sheet=load_workbook(before,read_only=True,keep_vba=True).sheetnames[0]
    diff=compare_cells(before,after,sheet); emit(diff[:500]); print(f'\nchanges on {sheet}: {len(diff)}')


def cmd_safe(path,sheet,edits,output,backup):
    emit(__import__('excel_power_engine.safe_edit', fromlist=['SafeEditor']).SafeEditor().edit_cells(path,sheet,parse_edits(edits),output,backup).as_dict())


def cmd_smart(path,sheet,edits,output,backup,recalculate,prefer_native):
    ops=parse_edits(edits)
    emit(ExcelEngine().smart_edit(path,sheet,ops,output,recalculate=recalculate,prefer_native=prefer_native,create_backup=backup))


def cmd_recalculate(path, output, backup, full_rebuild, visible):
    emit(smart_recalculate(path, output, create_backup=backup, full_rebuild=full_rebuild, visible=visible))


def cmd_open_excel(path, sheet, cell):
    from .open_excel import open_in_excel
    emit(open_in_excel(path, sheet, cell))


def cmd_cell_info(path, sheet, cell, trace=False, issues=False):
    engine=ExcelEngine()
    result={'cell':engine.cell_info(path,sheet,cell).as_dict()}
    if trace:
        result['trace']=engine.formula_trace(path,sheet,cell)
    if issues:
        result['formula_issues']=engine.formula_issues(path)
    emit(result)


def cmd_formula_scan(path):
    emit({'issues': ExcelEngine().formula_issues(path), 'issue_count': len(ExcelEngine().formula_issues(path))})


def cmd_bulk_preview(path, sheet, start_cell, data_file, data_sheet, no_formulas, clear_empty, ranges):
    from .bulk import preview_file, read_data_file, preview_multi_ranges
    matrix=read_data_file(data_file, sheet=data_sheet)
    if ranges:
        changes=preview_multi_ranges(path,sheet,ranges,matrix,interpret_formulas=not no_formulas,clear_empty=clear_empty)
        emit({"sheet":sheet,"target_ranges":ranges,"rows":len(matrix),"columns":max((len(r) for r in matrix),default=0),"changed_count":sum(c.before!=c.after for c in changes),"changes":[c.as_dict() for c in changes]})
    else:
        emit(preview_file(path,sheet,start_cell,data_file,data_sheet=data_sheet,interpret_formulas=not no_formulas,clear_empty=clear_empty))


def cmd_bulk_edit(path, sheet, start_cell, data_file, data_sheet, output, backup, no_formulas, clear_empty, ranges):
    from .bulk import execute_file, read_data_file, execute_multi_ranges
    matrix=read_data_file(data_file, sheet=data_sheet)
    if ranges:
        emit(execute_multi_ranges(path,sheet,ranges,matrix,output,interpret_formulas=not no_formulas,clear_empty=clear_empty,create_backup=backup))
    else:
        emit(execute_file(path,sheet,start_cell,data_file,output,data_sheet=data_sheet,interpret_formulas=not no_formulas,clear_empty=clear_empty,create_backup=backup))



def cmd_search(path, term, mode, case_sensitive, formulas, values, sheet, max_results):
    sheets=[sheet] if sheet else None
    hits=search_workbook(path, term, sheets=sheets, mode=mode, case_sensitive=case_sensitive, search_formulas=formulas, search_values=values, max_results=max_results)
    emit({"term":term,"mode":mode,"count":len(hits),"hits":[h.as_dict() for h in hits]})


def cmd_targets(path, sheet, specs, max_cells):
    emit({"sheet":sheet,"specs":[s.as_dict() for s in parse_target_specs(specs)],"cells":resolve_target_specs(path,sheet,specs,max_cells=max_cells)})


def cmd_native_format(path, sheet, ranges, output, args):
    kwargs={k:getattr(args,k) for k in ("fill","font_color","bold","italic","font_size","horizontal","vertical","wrap","number_format") if getattr(args,k) is not None}
    emit(ExcelOperations().format(path,sheet,ranges.split(";"),output=output,visible=args.visible,backup=not args.no_backup,**kwargs).as_dict())


def cmd_native_borders(path, sheet, ranges, output, args):
    emit(ExcelOperations().borders(path,sheet,ranges.split(";"),output=output,line_style=args.line_style,weight=args.weight,color=args.color,visible=args.visible,backup=not args.no_backup).as_dict())


def cmd_native_merge(path, sheet, ranges, output, args):
    emit(ExcelOperations().merge(path,sheet,ranges.split(";"),output=output,unmerge=args.unmerge,visible=args.visible,backup=not args.no_backup).as_dict())


def cmd_print_setup(path, sheet, output, args):
    emit(ExcelOperations().print_setup(path,sheet,output=output,print_area=args.print_area,orientation=args.orientation,paper_size=args.paper_size,fit_width=args.fit_width,fit_height=args.fit_height,repeat_rows=args.repeat_rows,repeat_columns=args.repeat_columns,visible=args.visible,backup=not args.no_backup).as_dict())


def cmd_export_pdf(path, output_pdf, sheet, ranges, visible):
    emit(ExcelOperations().export_pdf(path,output_pdf=output_pdf,sheet=sheet,ranges=ranges.split(";") if ranges else None,visible=visible))


def cmd_workflow_plan(path, workflow_file, output):
    ops=load_workflow(workflow_file)
    emit({"source":str(path),"output":str(output),"valid":True,"preview":preview_operations(validate_workflow(ops)),"operations":ops})


def cmd_workflow_run(path, workflow_file, output, visible, no_backup):
    emit(run_workflow(path,load_workflow(workflow_file),output,backup=not no_backup,visible=visible))

def cmd_workspace_plan(path, workflow_file):
    emit(preview_workspace(path, workflow_file))


def cmd_workspace_run(path, workflow_file, output, visible, no_backup, no_open):
    emit(run_workspace(path, workflow_file, output, backup=not no_backup, visible=visible, open_after=not no_open))


def cmd_workspace_rollback(output):
    emit(rollback_workspace(output))

def cmd_command_plan(path, command, sheet):
    from .workbook_context import analyze_workbook
    from .command_intelligence import build_command_plan
    ctx = analyze_workbook(path, selected_sheet=sheet)
    emit(build_command_plan(command, ctx, sheet=sheet).as_dict())


def main():
    p=argparse.ArgumentParser(prog='excel-power')
    sub=p.add_subparsers(dest='command',required=True)
    sub.add_parser('gui')

    a=sub.add_parser('inspect');a.add_argument('path')
    a=sub.add_parser('inspect-deep');a.add_argument('path')
    a=sub.add_parser('plan-edit');a.add_argument('path');a.add_argument('--operation',default='cell-edit');a.add_argument('--recalculate',action='store_true');a.add_argument('--prefer-native',action='store_true')
    a=sub.add_parser('read');a.add_argument('path');a.add_argument('--sheet',default='Sheet1');a.add_argument('--engine',default='auto')
    a=sub.add_parser('diff');a.add_argument('before');a.add_argument('after');a.add_argument('--sheet')
    a=sub.add_parser('safe-edit');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--edit',action='append',required=True);a.add_argument('--output');a.add_argument('--no-backup',action='store_true')
    a=sub.add_parser('smart-edit');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--edit',action='append',required=True);a.add_argument('--output');a.add_argument('--no-backup',action='store_true');a.add_argument('--recalculate',action='store_true');a.add_argument('--prefer-native',action='store_true')
    a=sub.add_parser('smart-recalc');a.add_argument('path');a.add_argument('--output');a.add_argument('--no-backup',action='store_true');a.add_argument('--no-full-rebuild',action='store_true');a.add_argument('--visible',action='store_true')
    a=sub.add_parser('open-excel');a.add_argument('path');a.add_argument('--sheet');a.add_argument('--cell')
    a=sub.add_parser('cell-info');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--cell',required=True);a.add_argument('--trace',action='store_true');a.add_argument('--issues',action='store_true')
    a=sub.add_parser('formula-scan');a.add_argument('path')
    a=sub.add_parser('bulk-preview');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--start-cell');a.add_argument('--ranges');a.add_argument('--data-file',required=True);a.add_argument('--data-sheet');a.add_argument('--no-formulas',action='store_true');a.add_argument('--clear-empty',action='store_true')
    a=sub.add_parser('bulk-edit');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--start-cell');a.add_argument('--ranges');a.add_argument('--data-file',required=True);a.add_argument('--data-sheet');a.add_argument('--output');a.add_argument('--no-backup',action='store_true');a.add_argument('--no-formulas',action='store_true');a.add_argument('--clear-empty',action='store_true')

    a=sub.add_parser('search');a.add_argument('path');a.add_argument('term');a.add_argument('--mode',choices=['contains','exact','starts','ends','regex'],default='contains');a.add_argument('--case-sensitive',action='store_true');a.add_argument('--no-formulas',action='store_true');a.add_argument('--no-values',action='store_true');a.add_argument('--sheet');a.add_argument('--max-results',type=int,default=2000)
    a=sub.add_parser('resolve-targets');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--targets',required=True);a.add_argument('--max-cells',type=int,default=200000)
    a=sub.add_parser('format');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--ranges',required=True);a.add_argument('--output');a.add_argument('--fill');a.add_argument('--font-color');a.add_argument('--bold',action=argparse.BooleanOptionalAction,default=None);a.add_argument('--italic',action=argparse.BooleanOptionalAction,default=None);a.add_argument('--font-size',type=float);a.add_argument('--horizontal',choices=['left','center','right']);a.add_argument('--vertical',choices=['top','middle','bottom']);a.add_argument('--wrap',action=argparse.BooleanOptionalAction,default=None);a.add_argument('--number-format');a.add_argument('--visible',action='store_true');a.add_argument('--no-backup',action='store_true')
    a=sub.add_parser('borders');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--ranges',required=True);a.add_argument('--output');a.add_argument('--line-style',type=int,default=1);a.add_argument('--weight',type=int,default=2);a.add_argument('--color',type=int);a.add_argument('--visible',action='store_true');a.add_argument('--no-backup',action='store_true')
    a=sub.add_parser('merge');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--ranges',required=True);a.add_argument('--output');a.add_argument('--unmerge',action='store_true');a.add_argument('--visible',action='store_true');a.add_argument('--no-backup',action='store_true')
    a=sub.add_parser('print-setup');a.add_argument('path');a.add_argument('--sheet',required=True);a.add_argument('--output');a.add_argument('--print-area');a.add_argument('--orientation',choices=['portrait','landscape']);a.add_argument('--paper-size',type=int);a.add_argument('--fit-width',type=int);a.add_argument('--fit-height',type=int);a.add_argument('--repeat-rows');a.add_argument('--repeat-columns');a.add_argument('--visible',action='store_true');a.add_argument('--no-backup',action='store_true')
    a=sub.add_parser('export-pdf');a.add_argument('path');a.add_argument('--output-pdf');a.add_argument('--sheet');a.add_argument('--ranges');a.add_argument('--visible',action='store_true')
    a=sub.add_parser('workflow-plan');a.add_argument('path');a.add_argument('--workflow',required=True);a.add_argument('--output',required=True)
    a=sub.add_parser('workflow-run');a.add_argument('path');a.add_argument('--workflow',required=True);a.add_argument('--output',required=True);a.add_argument('--visible',action='store_true');a.add_argument('--no-backup',action='store_true')

    a=sub.add_parser('workspace-plan');a.add_argument('path');a.add_argument('--workflow',required=True)
    a=sub.add_parser('workspace-run');a.add_argument('path');a.add_argument('--workflow',required=True);a.add_argument('--output',required=True);a.add_argument('--visible',action='store_true');a.add_argument('--no-backup',action='store_true');a.add_argument('--no-open',action='store_true')
    a=sub.add_parser('workspace-rollback');a.add_argument('output')
    a=sub.add_parser('command-plan');a.add_argument('path');a.add_argument('text');a.add_argument('--sheet')

    x=p.parse_args()
    if x.command=='gui':
        from .gui import main as gui_main
        return gui_main()
    if x.command=='inspect':cmd_inspect(x.path,False)
    elif x.command=='inspect-deep':cmd_inspect(x.path,True)
    elif x.command=='plan-edit':cmd_plan(x.path,x.operation,x.recalculate,x.prefer_native)
    elif x.command=='read':cmd_read(x.path,x.sheet,x.engine)
    elif x.command=='diff':cmd_diff(x.before,x.after,x.sheet)
    elif x.command=='safe-edit':cmd_safe(x.path,x.sheet,x.edit,x.output,not x.no_backup)
    elif x.command=='smart-edit':cmd_smart(x.path,x.sheet,x.edit,x.output,not x.no_backup,x.recalculate,x.prefer_native)
    elif x.command=='smart-recalc':cmd_recalculate(x.path,x.output,not x.no_backup,not x.no_full_rebuild,x.visible)
    elif x.command=='open-excel':cmd_open_excel(x.path,x.sheet,x.cell)
    elif x.command=='cell-info':cmd_cell_info(x.path,x.sheet,x.cell,x.trace,x.issues)
    elif x.command=='formula-scan':cmd_formula_scan(x.path)
    elif x.command=='bulk-preview':
        if not x.ranges and not x.start_cell: raise SystemExit('Provide --start-cell or --ranges')
        cmd_bulk_preview(x.path,x.sheet,x.start_cell,x.data_file,x.data_sheet,x.no_formulas,x.clear_empty,x.ranges)
    elif x.command=='bulk-edit':
        if not x.ranges and not x.start_cell: raise SystemExit('Provide --start-cell or --ranges')
        cmd_bulk_edit(x.path,x.sheet,x.start_cell,x.data_file,x.data_sheet,x.output,not x.no_backup,x.no_formulas,x.clear_empty,x.ranges)
    elif x.command=='search': cmd_search(x.path,x.term,x.mode,x.case_sensitive,not x.no_formulas,not x.no_values,x.sheet,x.max_results)
    elif x.command=='resolve-targets': cmd_targets(x.path,x.sheet,x.targets,x.max_cells)
    elif x.command=='format': cmd_native_format(x.path,x.sheet,x.ranges,x.output,x)
    elif x.command=='borders': cmd_native_borders(x.path,x.sheet,x.ranges,x.output,x)
    elif x.command=='merge': cmd_native_merge(x.path,x.sheet,x.ranges,x.output,x)
    elif x.command=='print-setup': cmd_print_setup(x.path,x.sheet,x.output,x)
    elif x.command=='export-pdf': cmd_export_pdf(x.path,x.output_pdf,x.sheet,x.ranges,x.visible)
    elif x.command=='workflow-plan': cmd_workflow_plan(x.path,x.workflow,x.output)
    elif x.command=='workflow-run': cmd_workflow_run(x.path,x.workflow,x.output,x.visible,x.no_backup)
    elif x.command=='workspace-plan': cmd_workspace_plan(x.path,x.workflow)
    elif x.command=='workspace-run': cmd_workspace_run(x.path,x.workflow,x.output,x.visible,x.no_backup,x.no_open)
    elif x.command=='workspace-rollback': cmd_workspace_rollback(x.output)
    elif x.command=='command-plan': cmd_command_plan(x.path,x.text,x.sheet)
    return 0

if __name__=='__main__':raise SystemExit(main())
