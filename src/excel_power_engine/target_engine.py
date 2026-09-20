from __future__ import annotations

import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable

from .sheet_viewer import workbook_sheets

CELL_RE = re.compile(r"^([A-Za-z]+)(\d+)$")
RANGE_RE = re.compile(r"^([A-Za-z]+\d+)\s*:\s*([A-Za-z]+\d+)$")
ROW_RE = re.compile(r"^(\d+)(?::(\d+))?$")
COL_RE = re.compile(r"^([A-Za-z]+)(?::([A-Za-z]+))?$")


def col_to_num(col: str) -> int:
    n = 0
    for ch in col.upper():
        if not ch.isalpha():
            raise ValueError(f"Invalid column: {col}")
        n = n * 26 + ord(ch) - 64
    return n


def num_to_col(n: int) -> str:
    if n < 1:
        raise ValueError("Column number must be >= 1")
    out = ""
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def cell_to_rc(cell: str) -> tuple[int, int]:
    m = CELL_RE.fullmatch(cell.strip())
    if not m:
        raise ValueError(f"Invalid cell reference: {cell}")
    return int(m.group(2)), col_to_num(m.group(1))


@dataclass(frozen=True)
class TargetSpec:
    raw: str
    kind: str
    start_row: int | None = None
    start_col: int | None = None
    end_row: int | None = None
    end_col: int | None = None

    def as_dict(self) -> dict:
        return asdict(self)


def parse_target_specs(text: str) -> list[TargetSpec]:
    if not text or not text.strip():
        return []
    out: list[TargetSpec] = []
    for token in re.split(r"[;,]+", text):
        token = token.strip()
        if not token:
            continue
        m = RANGE_RE.fullmatch(token)
        if m:
            sr, sc = cell_to_rc(m.group(1))
            er, ec = cell_to_rc(m.group(2))
            out.append(TargetSpec(token, "range", min(sr, er), min(sc, ec), max(sr, er), max(sc, ec)))
            continue
        m = CELL_RE.fullmatch(token)
        if m:
            r, c = cell_to_rc(token)
            out.append(TargetSpec(token, "cell", r, c, r, c))
            continue
        m = ROW_RE.fullmatch(token)
        if m:
            sr = int(m.group(1)); er = int(m.group(2) or sr)
            out.append(TargetSpec(token, "row", sr, None, max(sr, er), None))
            continue
        m = COL_RE.fullmatch(token)
        if m:
            sc = col_to_num(m.group(1)); ec = col_to_num(m.group(2) or m.group(1))
            out.append(TargetSpec(token, "column", None, min(sc, ec), None, max(sc, ec)))
            continue
        raise ValueError(f"Invalid target specification: {token}")
    return out


def resolve_target_specs(path: str | Path, sheet: str, text: str, *, max_cells: int = 200_000) -> list[str]:
    specs = parse_target_specs(text)
    sheets = workbook_sheets(path)
    if sheet not in sheets:
        raise KeyError(f"Worksheet not found: {sheet}")
    # Read the used dimensions with openpyxl only for discovery; writes remain in the engine.
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=True)
    ws = wb[sheet]
    max_row = max(1, ws.max_row)
    max_col = max(1, ws.max_column)
    out: list[str] = []
    for spec in specs:
        sr = spec.start_row if spec.start_row is not None else 1
        er = spec.end_row if spec.end_row is not None else max_row
        sc = spec.start_col if spec.start_col is not None else 1
        ec = spec.end_col if spec.end_col is not None else max_col
        er = min(er, max_row); ec = min(ec, max_col)
        for r in range(sr, er + 1):
            for c in range(sc, ec + 1):
                out.append(f"{num_to_col(c)}{r}")
                if len(out) >= max_cells:
                    raise ValueError(f"Target expansion exceeds max_cells={max_cells}")
    return out
