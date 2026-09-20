from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from excel_power_engine.smart_engine import build_plan, SmartEditEngine
from excel_power_engine.safe_edit import EditOperation


def make_simple(path: Path):
    import openpyxl
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Data'; ws['A1']='Name'; ws['B1']='Value'; ws['A2']='Alpha'; ws['B2']=10; wb.save(path)


def inject_sensitive_xlsm(src: Path, dst: Path):
    with ZipFile(src,'r') as zin, ZipFile(dst,'w',ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            data=zin.read(i.filename)
            if i.filename=='xl/worksheets/sheet1.xml':
                txt=data.decode()
                ext='<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations count="1"><x14:dataValidation type="list"><x14:formula1><xm:f>$J$2:$J$14</xm:f></x14:formula1><xm:sqref>B2:B3</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>'
                data=txt.replace('</worksheet>',ext+'</worksheet>').encode()
            zout.writestr(i,data)
        zout.writestr('xl/vbaProject.bin',b'FAKE-VBA')


def test_plan_simple_prefers_openpyxl(tmp_path):
    p=tmp_path/'simple.xlsx'; make_simple(p)
    plan=build_plan(p,'cell-edit')
    assert plan.engine=='openpyxl'
    assert plan.safety=='standard'


def test_plan_sensitive_prefers_surgical(tmp_path):
    base=tmp_path/'base.xlsx'; make_simple(base)
    p=tmp_path/'macro.xlsm'; inject_sensitive_xlsm(base,p)
    plan=build_plan(p,'cell-edit')
    assert plan.engine=='ooxml-surgical'
    assert 'x14' in plan.sensitive_features
    assert plan.has_vba


def test_smart_edit_simple_and_preserve_sensitive(tmp_path):
    p=tmp_path/'simple.xlsx'; make_simple(p)
    out=tmp_path/'simple_out.xlsx'
    result=SmartEditEngine().edit_cells(p,'Data',[EditOperation('B2',99)],out,create_backup=True)
    assert out.exists() and result['executed_engine']=='openpyxl'

    base=tmp_path/'base2.xlsx'; make_simple(base)
    sensitive=tmp_path/'macro.xlsm'; inject_sensitive_xlsm(base,sensitive)
    out2=tmp_path/'macro_out.xlsm'
    result2=SmartEditEngine().edit_cells(sensitive,'Data',[EditOperation('B2',77)],out2,create_backup=True)
    assert out2.exists() and result2['executed_engine']=='ooxml-surgical'
    assert result2['vba_preserved'] and result2['x14_preserved']
