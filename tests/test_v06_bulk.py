from pathlib import Path
import csv

from excel_power_engine.bulk import parse_matrix, preview_matrix, execute_matrix, read_data_file, parse_target_ranges, preview_multi_ranges, execute_multi_ranges
from excel_power_engine.formula_intelligence import cell_info


def make_book(path: Path):
    import openpyxl
    wb=openpyxl.Workbook()
    ws=wb.active; ws.title='الشهر الخامس'
    ws['F11']='old'; ws['G11']=0; ws['H11']='old'; ws['I11']=0
    ws['F12']='old2'; ws['G12']=1; ws['H12']='old2'; ws['I12']=1
    wb.save(path)


def test_parse_and_preview_bulk_matrix(tmp_path):
    p=tmp_path/'demo.xlsx'; make_book(p)
    matrix=parse_matrix('2026/05/01\t3\tالتهاب الضرع\t4\n2026/05/02\t2\tطفيليات داخلية\t5')
    changes=preview_matrix(p,'الشهر الخامس','F11',matrix)
    assert len(changes)==8
    assert any(c.cell=='G11' and c.after==3 for c in changes)


def test_execute_bulk_preserves_cell_updates(tmp_path):
    p=tmp_path/'demo.xlsx'; make_book(p)
    out=tmp_path/'out.xlsx'
    matrix=parse_matrix('2026/05/01\t3\tالتهاب الضرع\t4\n2026/05/02\t2\tطفيليات داخلية\t5')
    result=execute_matrix(p,'الشهر الخامس','F11',matrix,out)
    assert out.exists()
    assert result['bulk'] is True
    assert result['changed_cell_count']==8
    assert cell_info(out,'الشهر الخامس','G11').value == 3


def test_bulk_xlsm_preserves_vba_and_x14(tmp_path):
    from zipfile import ZipFile, ZIP_DEFLATED
    base=tmp_path/'base.xlsx'
    make_book(base)
    sensitive=tmp_path/'macro.xlsm'
    with ZipFile(base,'r') as zin, ZipFile(sensitive,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                text=data.decode()
                ext='<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations count="1"><x14:dataValidation type="list"><x14:formula1><xm:f>$J$2:$J$14</xm:f></x14:formula1><xm:sqref>F11:I12</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>'
                data=text.replace('</worksheet>', ext+'</worksheet>').encode()
            zout.writestr(item,data)
        zout.writestr('xl/vbaProject.bin', b'FAKE-VBA')
    out=tmp_path/'bulk_macro.xlsm'
    matrix=parse_matrix('2026/05/01\t3\tالتهاب الضرع\t4\n2026/05/02\t2\tطفيليات داخلية\t5')
    result=execute_matrix(sensitive,'الشهر الخامس','F11',matrix,out)
    assert result['vba_preserved']
    assert result['x14_preserved']


def test_multi_range_mapping(tmp_path):
    p=tmp_path/'demo.xlsx'; make_book(p)
    matrix=parse_matrix('A\t1\nB\t4')
    ranges=parse_target_ranges('F11:G11;F12:G12')
    assert len(ranges)==2
    changes=preview_multi_ranges(p,'الشهر الخامس','F11:G11;F12:G12',matrix)
    assert changes[0].cell=='F11'
    assert changes[1].cell=='G11'
    assert changes[2].cell=='F12'
    out=tmp_path/'multi.xlsx'
    result=execute_multi_ranges(p,'الشهر الخامس','F11:G11;F12:G12',matrix,out)
    assert result['changed_cell_count']==4
    assert cell_info(out,'الشهر الخامس','G12').value==4
