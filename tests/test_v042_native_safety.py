from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from excel_power_engine.native_excel import NativeExcelBridge


def make_xlsm(path: Path, vba: bytes, formula_error: bool = False):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = "=SUM(A1:A2)"
    if formula_error:
        ws["B1"] = "=1/0"
    tmp = path.with_suffix('.xlsx')
    wb.save(tmp)
    with ZipFile(tmp, 'r') as zin, ZipFile(path, 'w', ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            zout.writestr(info, data)
        zout.writestr("xl/vbaProject.bin", vba)
    tmp.unlink()


def test_preserve_vba_binary(tmp_path):
    src = tmp_path / "src.xlsm"
    out = tmp_path / "out.xlsm"
    make_xlsm(src, b"ORIGINAL-VBA")
    make_xlsm(out, b"CHANGED-VBA")
    original = NativeExcelBridge()._sha_vba if False else None
    from excel_power_engine.inspect import inspect_workbook
    before = inspect_workbook(src, deep=True)
    NativeExcelBridge._preserve_vba_binary(src, out, before.integrity["vba_sha256"])
    after = inspect_workbook(out, deep=True)
    assert after.integrity["vba_sha256"] == before.integrity["vba_sha256"]


def test_formula_error_scan_read_only(tmp_path):
    p = tmp_path / "err.xlsx"
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=1/0"
    wb.save(p)
    errors = NativeExcelBridge._scan_formula_errors(p)
    assert errors == []  # formula caches are not populated by openpyxl; no false positives
