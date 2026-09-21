from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .formula_intelligence import extract_references


def inspect_cell(path: str | Path, sheet: str, cell: str) -> dict[str, Any]:
    p = Path(path).resolve()
    wb = load_workbook(p, read_only=False, data_only=False, keep_vba=p.suffix.lower() == ".xlsm")
    try:
        ws = wb[sheet]
        c = ws[cell]
        style = c._style
        color = c.font.color
        fill_color = c.fill.fgColor.rgb or c.fill.fgColor.indexed or c.fill.fgColor.theme
        dep = []
        if isinstance(c.value, str) and c.value.startswith("="):
            try:
                dep = sorted(extract_references(c.value))
            except Exception:
                dep = []
        merged = any(c.coordinate in r for r in ws.merged_cells.ranges)
        return {
            "sheet": sheet,
            "cell": c.coordinate,
            "value": c.value,
            "formula": c.value if isinstance(c.value, str) and c.value.startswith("=") else None,
            "data_type": c.data_type,
            "number_format": c.number_format,
            "font": {"name": c.font.name, "size": c.font.sz, "bold": c.font.bold, "italic": c.font.italic, "underline": c.font.underline,
                     "color": color.type + ":" + str(color.rgb if color.type == "rgb" else color.indexed if color.type == "indexed" else color.theme) if color and color.type else None},
            "fill": {"fill_type": c.fill.fill_type, "color": fill_color},
            "alignment": {"horizontal": c.alignment.horizontal, "vertical": c.alignment.vertical, "wrap_text": c.alignment.wrap_text},
            "borders": {side: getattr(c.border, side).style for side in ("left", "right", "top", "bottom")},
            "merged": merged,
            "row_height": ws.row_dimensions[c.row].height,
            "column_width": ws.column_dimensions[c.column_letter].width,
            "sheet_hidden": ws.sheet_state,
            "dependencies": dep,
        }
    finally:
        wb.close()
