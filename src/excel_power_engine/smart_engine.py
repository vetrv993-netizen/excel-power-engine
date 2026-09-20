from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from .inspect import inspect_workbook

SENSITIVE_FEATURES = {
    "vba",
    "x14",
    "controls",
    "slicers",
    "vml",
    "pivot",
    "customXml",
    "external-links",
    "connections",
    "query-tables",
}

@dataclass(slots=True)
class SmartPlan:
    path: Path
    operation: str
    engine: str
    reason: str
    safety: str
    is_ooxml: bool
    is_xlsm: bool
    has_vba: bool
    sensitive_features: list[str] = field(default_factory=list)
    native_excel_available: bool = False
    requested_recalculate: bool = False
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": str(self.path),
            "operation": self.operation,
            "selected_engine": self.engine,
            "reason": self.reason,
            "safety": self.safety,
            "is_ooxml": self.is_ooxml,
            "is_xlsm": self.is_xlsm,
            "has_vba": self.has_vba,
            "sensitive_features": self.sensitive_features,
            "native_excel_available": self.native_excel_available,
            "requested_recalculate": self.requested_recalculate,
            "warnings": self.warnings,
        }


def _native_excel_available() -> bool:
    # Fast detection without importing Windows-only modules on other platforms.
    import sys
    if sys.platform != "win32":
        return False
    try:
        import xlwings  # noqa: F401
        return True
    except Exception:
        try:
            import win32com.client  # noqa: F401
            return True
        except Exception:
            return False


def build_plan(path: str | Path, operation: str, *, recalculate: bool = False, prefer_native: bool = False) -> SmartPlan:
    p = Path(path)
    profile = inspect_workbook(p, deep=True)
    op = operation.strip().lower().replace("_", "-")
    native = _native_excel_available()
    sensitive = sorted(f for f in profile.features if f in SENSITIVE_FEATURES)
    has_sensitive_xml = bool(sensitive)
    warnings: list[str] = []

    if op in {"inspect", "inspect-deep", "profile"}:
        return SmartPlan(p, op, "ooxml-inspector", "الفحص لا يحتاج إلى إعادة حفظ المصنف.", "read-only", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
    if op in {"create", "report"}:
        return SmartPlan(p, op, "xlsxwriter", "العملية إنشاء مصنف جديد، وليست تحرير ملف موجود.", "safe", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
    if op in {"vba-inspect", "vba"}:
        return SmartPlan(p, op, "oletools", "قراءة/تحليل VBA يحتاج محلل OLE/VBA وليس محرر خلايا.", "read-only", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
    if op in {"read-fast", "analyze", "analysis"}:
        return SmartPlan(p, op, "fast-reader", "العملية قراءة وتحليل؛ لا حاجة لإعادة حفظ المصنف.", "read-only", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
    if op in {"xml-edit", "surgical-edit", "safe-edit"}:
        return SmartPlan(p, op, "ooxml-surgical", "اختيار صريح للتحرير الجراحي على مستوى OOXML.", "preserve", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
    if op in {"recalculate", "recalc"} or recalculate:
        if native:
            return SmartPlan(p, op, "excel-native", "إعادة الحساب الفعلية تعتمد على Excel Object Model.", "native", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, True, warnings)
        warnings.append("Excel Native غير متاح؛ لن يتم ضمان إعادة حساب صيغ Excel كما يفعل Excel الحقيقي.")
        if profile.is_xlsm or has_sensitive_xml:
            return SmartPlan(p, op, "ooxml-surgical", "Excel غير متاح، والحفاظ على أجزاء XLSM الحساسة أهم من إعادة الحفظ عبر OpenPyXL.", "preserve-with-warning", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, True, warnings)
        return SmartPlan(p, op, "openpyxl", "Excel غير متاح والملف لا يحتوي امتدادات حساسة؛ يمكن حفظه عبر OpenPyXL مع تحذير أن إعادة الحساب ليست مضمونة.", "warning", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, True, warnings)
    if op in {"cell-edit", "formula-edit", "edit", "format-edit"}:
        # A native Excel request may be preferable on Windows for operations that need Excel's object model.
        if prefer_native and native:
            return SmartPlan(p, op, "excel-native", "تم تفعيل تفضيل Excel Native وتوفر Excel على Windows.", "native", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
        if profile.is_xlsm or has_sensitive_xml:
            reason_bits = []
            if profile.is_xlsm:
                reason_bits.append("XLSM/VBA")
            if has_sensitive_xml:
                reason_bits.append("امتدادات حساسة: " + ", ".join(sensitive))
            return SmartPlan(p, op, "ooxml-surgical", " و".join(reason_bits) + "; التحرير الجراحي يقلل إعادة كتابة الحزمة.", "preserve", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
        return SmartPlan(p, op, "openpyxl", "ملف OOXML بسيط بلا امتدادات حساسة؛ OpenPyXL مناسب للتحرير الهيكلي.", "standard", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)
    return SmartPlan(p, op, "ooxml-inspector", "عملية غير معروفة؛ يبدأ المحرك بالفحص بدلاً من تعديل الملف تخمينيًا.", "read-only", profile.is_ooxml, profile.is_xlsm, profile.has_vba, sensitive, native, recalculate, warnings)


class SmartEditEngine:
    def plan(self, path: str | Path, operation: str = "cell-edit", *, recalculate: bool = False, prefer_native: bool = False) -> SmartPlan:
        return build_plan(path, operation, recalculate=recalculate, prefer_native=prefer_native)

    def edit_cells(
        self,
        source: str | Path,
        sheet: str,
        edits: Iterable[Any],
        output: str | Path | None = None,
        *,
        recalculate: bool = False,
        prefer_native: bool = False,
        create_backup: bool = True,
    ) -> dict[str, Any]:
        ops = list(edits)
        plan = self.plan(source, "cell-edit", recalculate=recalculate, prefer_native=prefer_native)
        if plan.engine == "ooxml-surgical":
            from .safe_edit import SafeEditor
            report = SafeEditor().edit_cells(source, sheet, ops, output, create_backup)
            result = report.as_dict()
            result.update({"smart_plan": plan.as_dict(), "executed_engine": plan.engine})
            if recalculate:
                result.setdefault("warnings", []).extend(plan.warnings or ["لم يتم تنفيذ Recalculate عبر Excel Native."])
            return result
        if plan.engine == "openpyxl":
            return _openpyxl_edit(source, sheet, ops, output, create_backup, plan)
        if plan.engine == "excel-native":
            from .native_excel import NativeExcelEditor
            return NativeExcelEditor().edit_cells(source, sheet, ops, output, create_backup=create_backup, recalculate=recalculate, plan=plan)
        raise RuntimeError(f"Smart Engine selected unsupported execution engine: {plan.engine}")


def _openpyxl_edit(source, sheet, edits, output, create_backup, plan: SmartPlan) -> dict[str, Any]:
    from shutil import copy2
    from openpyxl import load_workbook
    source = Path(source)
    output = Path(output) if output else source.with_name(source.stem + "_SMART_EDIT" + source.suffix)
    if source.resolve() == output.resolve():
        raise ValueError("SmartEdit requires a different output path")
    backup = source.with_name(source.stem + "_BACKUP" + source.suffix) if create_backup else None
    if backup:
        copy2(source, backup)
    wb = load_workbook(source, keep_vba=False, data_only=False)
    if sheet not in wb.sheetnames:
        raise KeyError(f"Worksheet not found: {sheet}")
    ws = wb[sheet]
    changed = []
    for op in edits:
        cell = op.cell.upper()
        ws[cell] = op.formula if op.formula is not None else op.value
        changed.append(cell)
    wb.save(output)
    return {
        "source": str(source), "output": str(output), "backup": str(backup) if backup else None,
        "sheet": sheet, "edited_cells": changed, "changed_parts": [f"worksheet:{sheet}"],
        "protected_parts_unchanged": True, "vba_preserved": False, "x14_preserved": False,
        "warnings": ["OpenPyXL was selected only because the file was classified as a simple workbook."],
        "smart_plan": plan.as_dict(), "executed_engine": "openpyxl",
    }
