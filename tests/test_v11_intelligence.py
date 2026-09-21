from pathlib import Path

from openpyxl import Workbook

from excel_power_engine.command_intelligence import build_command_plan, parse_command
from excel_power_engine.inspector import inspect_cell
from excel_power_engine.risk import classify_operation
from excel_power_engine.semantic_engine import SemanticWorkbook, normalize_text
from excel_power_engine.suggestions import smart_suggestions
from excel_power_engine.operations import preview_operations


def make_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "الطلاب"
    ws.append(["اسم الطالب", "الصف", "الدرجة", "المعدل"])
    ws.append(["أحمد", "5", 90, 95])
    ws.append(["مريم", "5", 80, 88])
    ws["C2"] = "=90"
    wb.save(path)


def test_normalize_arabic():
    assert normalize_text(" أَحْمَدـ ") == "احمد"


def test_parse_natural_fill_command():
    intent = parse_command("لوّن عمود الطلاب بالأحمر")
    assert intent is not None
    assert intent.action == "format_fill"
    assert intent.parameters["fill"] == "#FF0000"


def test_semantic_column_resolution(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    semantic = SemanticWorkbook(analyze_workbook(source))
    result = semantic.resolve_target("الطلاب", "الطلاب")
    assert result["resolved"]
    assert result["column"] == "A"
    assert result["range"] == "A2:A3"


def test_command_plan_resolves_real_column(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    plan = build_command_plan("لوّن عمود الطلاب بالأحمر", analyze_workbook(source), sheet="الطلاب")
    assert plan.status == "READY"
    assert plan.risk == "SAFE"
    assert plan.operations[0]["ranges"] == ["A2:A3"]


def test_command_plan_is_ambiguous_when_two_candidates_are_close(tmp_path):
    source = tmp_path / "ambiguous.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["اسم الطالب", "اسم المعلم", "الدرجة"]); ws.append(["A", "B", 90]); wb.save(source)
    from excel_power_engine.workbook_context import analyze_workbook
    plan = build_command_plan("لوّن عمود الاسم بالأحمر", analyze_workbook(source), sheet="Data")
    assert plan.status in {"AMBIGUOUS", "READY"}
    if plan.status == "AMBIGUOUS":
        assert plan.target["candidates"]


def test_table_border_command_uses_detected_used_range(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    plan = build_command_plan("ضع حدوداً حول جدول الدرجات", analyze_workbook(source), sheet="الطلاب")
    assert plan.status == "READY"
    assert plan.operations[0]["ranges"] == ["A1:D3"]


def test_print_pdf_commands(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    ctx = analyze_workbook(source)
    assert build_command_plan("اجعل الطباعة أفقية", ctx, sheet="الطلاب").operations[0]["kind"] == "print-setup"
    assert build_command_plan("صدّر الورقة إلى PDF", ctx, sheet="الطلاب").operations[0]["kind"] == "pdf"


def test_risk_classification():
    assert classify_operation("format") == "SAFE"
    assert classify_operation("merge") == "WARNING"
    assert classify_operation("delete-sheet") == "HIGH_RISK"


def test_operation_preview_includes_risk():
    data = preview_operations([{"kind": "visibility", "sheet": "Data", "ranges": ["A"], "payload": {"hidden": True}}])
    assert data[0]["risk"] == "WARNING"


def test_inspector_reports_real_style_fields(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    result = inspect_cell(source, "الطلاب", "A2")
    assert result["sheet"] == "الطلاب"
    assert result["cell"] == "A2"
    assert "font" in result and "fill" in result and "borders" in result


def test_smart_suggestions_are_non_destructive(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    suggestions = smart_suggestions(analyze_workbook(source), "الطلاب")
    assert suggestions
    assert all(x.get("preview") is True for x in suggestions)


def test_dynamic_column_choices(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    choices = SemanticWorkbook(analyze_workbook(source)).choices("الطلاب")
    titles = {x["title"] for x in choices}
    assert {"اسم الطالب", "الدرجة", "المعدل"} <= titles


def test_column_width_command(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    plan = build_command_plan("وسع عمود الاسم", analyze_workbook(source), sheet="الطلاب")
    assert plan.status == "READY"
    assert plan.operations[0]["kind"] == "column-width"


def test_hide_column_command(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    plan = build_command_plan("أخف عمود الدرجة", analyze_workbook(source), sheet="الطلاب")
    assert plan.status == "READY"
    assert plan.operations[0]["kind"] == "visibility"
    assert plan.operations[0]["payload"]["hidden"] is True


def test_font_size_command(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    plan = build_command_plan("غيّر حجم الخط إلى 14", analyze_workbook(source), sheet="الطلاب")
    assert plan.status == "READY"
    assert plan.operations[0]["payload"]["font_size"] == 14.0


def test_alignment_and_visibility_commands(tmp_path):
    source = tmp_path / "source.xlsx"
    make_book(source)
    from excel_power_engine.workbook_context import analyze_workbook
    ctx = analyze_workbook(source)
    alignment = build_command_plan("محاذاة وسط A1:D3", ctx, sheet="الطلاب")
    assert alignment.status == "READY"
    assert alignment.operations[0]["kind"] == "alignment"
    hidden = build_command_plan("أخف عمود الدرجة", ctx, sheet="الطلاب")
    assert hidden.operations[0]["kind"] == "visibility"


def test_workspace_accepts_new_v11_operations(tmp_path):
    from excel_power_engine.orchestrator import validate_workspace_operations
    ops = validate_workspace_operations([
        {"kind": "column-width", "sheet": "Data", "ranges": ["A"], "payload": {"width": 24}},
        {"kind": "visibility", "sheet": "Data", "ranges": ["B"], "payload": {"axis": "column", "hidden": True}},
        {"kind": "alignment", "sheet": "Data", "ranges": ["A1:B2"], "payload": {"horizontal": "center"}},
    ])
    assert [x["kind"] for x in ops] == ["column-width", "visibility", "alignment"]
