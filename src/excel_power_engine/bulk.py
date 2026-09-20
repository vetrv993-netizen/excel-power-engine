from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
import csv
import io
import json
import re
from typing import Any, Iterable

from .formula_intelligence import cell_info
from .safe_edit import EditOperation, SafeEditor
from .sheet_viewer import col_to_num, num_to_col, workbook_sheets

CELL_RE = re.compile(r"^([A-Za-z]+)(\d+)$")

@dataclass(slots=True)
class BulkChange:
    sheet: str
    cell: str
    before: Any
    after: Any
    kind: str = "value"

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def split_start_cell(ref: str) -> tuple[int, int]:
    m = CELL_RE.fullmatch(ref.strip())
    if not m:
        raise ValueError(f"Invalid start cell: {ref}")
    return int(m.group(2)), col_to_num(m.group(1).upper())


def parse_matrix(text: str, delimiter: str = "auto") -> list[list[str]]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line for line in text.split("\n") if line.strip()]
    if not lines:
        return []
    if delimiter == "auto":
        candidates = ["\t", ",", ";", "|"]
        delimiter = max(candidates, key=lambda d: max((line.count(d) for line in lines), default=0))
        if all(line.count(delimiter) == 0 for line in lines):
            delimiter = "\t"
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    return [list(row) for row in reader]


def read_data_file(path: str | Path, sheet: str | None = None) -> list[list[str]]:
    p = Path(path)
    if p.suffix.lower() in {".csv", ".txt", ".tsv"}:
        delimiter = "\t" if p.suffix.lower() == ".tsv" else ","
        return parse_matrix(p.read_text(encoding="utf-8-sig"), delimiter)
    if p.suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=False, keep_vba=True)
        ws = wb[sheet or wb.sheetnames[0]]
        out: list[list[str]] = []
        for row in ws.iter_rows(values_only=False):
            values = []
            for c in row:
                if c.value is None:
                    values.append("")
                else:
                    values.append(str(c.value))
            while values and values[-1] == "":
                values.pop()
            if values:
                out.append(values)
        return out
    raise ValueError(f"Unsupported bulk data file: {p.suffix}")


def _coerce(value: str, *, interpret_formulas: bool = True) -> tuple[Any, str | None]:
    if value is None:
        return None, None
    s = str(value)
    if s == "":
        return None, None
    if interpret_formulas and s.startswith("="):
        return None, s
    low = s.strip().lower()
    if low in {"true", "false"}:
        return low == "true", None
    if re.fullmatch(r"-?\d+", s.strip()):
        try:
            return int(s.strip()), None
        except ValueError:
            pass
    if re.fullmatch(r"-?(?:\d+\.\d*|\d*\.\d+)", s.strip()):
        try:
            return float(s.strip()), None
        except ValueError:
            pass
    return s, None


def matrix_edits(sheet: str, start_cell: str, matrix: list[list[str]], *, interpret_formulas: bool = True, clear_empty: bool = False) -> list[EditOperation]:
    start_row, start_col = split_start_cell(start_cell)
    ops: list[EditOperation] = []
    width = max((len(row) for row in matrix), default=0)
    for r_idx, row in enumerate(matrix):
        for c_idx in range(width):
            raw = row[c_idx] if c_idx < len(row) else ""
            if raw == "" and not clear_empty:
                continue
            value, formula = _coerce(raw, interpret_formulas=interpret_formulas)
            cell = f"{num_to_col(start_col + c_idx)}{start_row + r_idx}"
            ops.append(EditOperation(cell, value=value, formula=formula))
    return ops



def parse_target_ranges(text: str) -> list[tuple[int, int, int, int]]:
    """Parse semicolon/comma-separated rectangular Excel ranges."""
    if not text or not text.strip():
        return []
    out = []
    for raw in re.split(r"[;,]+", text):
        token = raw.strip()
        m = re.fullmatch(r"([A-Za-z]+\d+)\s*:\s*([A-Za-z]+\d+)", token)
        if not m:
            raise ValueError(f"Invalid target range: {token}")
        sr, sc = split_start_cell(m.group(1))
        er, ec = split_start_cell(m.group(2))
        out.append((min(sr, er), min(sc, ec), max(sr, er), max(sc, ec)))
    widths = {ec - sc + 1 for _, sc, _, ec in out}
    if len(widths) > 1:
        raise ValueError("All target ranges must have the same number of columns for bulk row mapping.")
    return out


def multi_range_edits(sheet: str, ranges_text: str, matrix: list[list[str]], *, interpret_formulas: bool = True, clear_empty: bool = False) -> list[EditOperation]:
    ranges = parse_target_ranges(ranges_text)
    if not ranges:
        raise ValueError("No target ranges supplied")
    range_width = ranges[0][3] - ranges[0][1] + 1
    total_rows = sum(r[2] - r[0] + 1 for r in ranges)
    if len(matrix) > total_rows:
        raise ValueError(f"Input contains {len(matrix)} rows but target ranges provide only {total_rows} rows")
    ops: list[EditOperation] = []
    row_cursor = 0
    for min_row, min_col, max_row, _ in ranges:
        for row_num in range(min_row, max_row + 1):
            if row_cursor >= len(matrix):
                break
            row = matrix[row_cursor]
            for col_offset in range(range_width):
                raw = row[col_offset] if col_offset < len(row) else ""
                if raw == "" and not clear_empty:
                    continue
                value, formula = _coerce(raw, interpret_formulas=interpret_formulas)
                cell = f"{num_to_col(min_col + col_offset)}{row_num}"
                ops.append(EditOperation(cell, value=value, formula=formula))
            row_cursor += 1
    return ops


def preview_multi_ranges(path: str | Path, sheet: str, ranges_text: str, matrix: list[list[str]], *, interpret_formulas: bool = True, clear_empty: bool = False) -> list[BulkChange]:
    ops = multi_range_edits(sheet, ranges_text, matrix, interpret_formulas=interpret_formulas, clear_empty=clear_empty)
    changes: list[BulkChange] = []
    for op in ops:
        snap = cell_info(path, sheet, op.cell)
        after = op.formula if op.formula is not None else op.value
        changes.append(BulkChange(sheet, op.cell, snap.formula if snap.formula is not None else snap.value, after, "formula" if op.formula is not None else "value"))
    return changes


def execute_multi_ranges(path: str | Path, sheet: str, ranges_text: str, matrix: list[list[str]], output: str | Path | None = None, *, interpret_formulas: bool = True, clear_empty: bool = False, create_backup: bool = True) -> dict[str, Any]:
    ops = multi_range_edits(sheet, ranges_text, matrix, interpret_formulas=interpret_formulas, clear_empty=clear_empty)
    report = SafeEditor().edit_cells(path, sheet, ops, output, create_backup=create_backup)
    result = report.as_dict()
    result.update({"bulk": True, "target_ranges": ranges_text, "rows": len(matrix), "columns": max((len(r) for r in matrix), default=0), "changed_cell_count": len(report.edited_cells), "clear_empty": clear_empty})
    return result

def preview_matrix(path: str | Path, sheet: str, start_cell: str, matrix: list[list[str]], *, interpret_formulas: bool = True, clear_empty: bool = False) -> list[BulkChange]:
    ops = matrix_edits(sheet, start_cell, matrix, interpret_formulas=interpret_formulas, clear_empty=clear_empty)
    changes: list[BulkChange] = []
    for op in ops:
        snap = cell_info(path, sheet, op.cell)
        after = op.formula if op.formula is not None else op.value
        kind = "formula" if op.formula is not None else "value"
        changes.append(BulkChange(sheet, op.cell, snap.formula if snap.formula is not None else snap.value, after, kind))
    return changes


def preview_file(path: str | Path, sheet: str, start_cell: str, data_file: str | Path, *, data_sheet: str | None = None, interpret_formulas: bool = True, clear_empty: bool = False) -> dict[str, Any]:
    matrix = read_data_file(data_file, sheet=data_sheet)
    changes = preview_matrix(path, sheet, start_cell, matrix, interpret_formulas=interpret_formulas, clear_empty=clear_empty)
    return {
        "sheet": sheet,
        "start_cell": start_cell.upper(),
        "rows": len(matrix),
        "columns": max((len(r) for r in matrix), default=0),
        "changes": [c.as_dict() for c in changes],
        "changed_count": sum(c.before != c.after for c in changes),
        "clear_empty": clear_empty,
    }


def execute_matrix(path: str | Path, sheet: str, start_cell: str, matrix: list[list[str]], output: str | Path | None = None, *, interpret_formulas: bool = True, clear_empty: bool = False, create_backup: bool = True) -> dict[str, Any]:
    ops = matrix_edits(sheet, start_cell, matrix, interpret_formulas=interpret_formulas, clear_empty=clear_empty)
    report = SafeEditor().edit_cells(path, sheet, ops, output, create_backup=create_backup)
    result = report.as_dict()
    result.update({
        "bulk": True,
        "start_cell": start_cell.upper(),
        "rows": len(matrix),
        "columns": max((len(r) for r in matrix), default=0),
        "changed_cell_count": len(report.edited_cells),
        "clear_empty": clear_empty,
    })
    return result


def execute_file(path: str | Path, sheet: str, start_cell: str, data_file: str | Path, output: str | Path | None = None, *, data_sheet: str | None = None, interpret_formulas: bool = True, clear_empty: bool = False, create_backup: bool = True) -> dict[str, Any]:
    matrix = read_data_file(data_file, sheet=data_sheet)
    return execute_matrix(path, sheet, start_cell, matrix, output, interpret_formulas=interpret_formulas, clear_empty=clear_empty, create_backup=create_backup)


def keyed_preview(path: str | Path, target_sheet: str, target_key_column: str, source_rows: list[dict[str, Any]], source_key_field: str, mapping: dict[str, str]) -> dict[str, Any]:
    """Map source records to target rows by a key column.

    mapping is {target_column: source_field}, e.g. {"F": "التاريخ", "G": "عدد المرضى"}.
    """
    p = Path(path)
    if target_sheet not in workbook_sheets(p):
        raise KeyError(f"Worksheet not found: {target_sheet}")
    # Read target key column row-by-row until the used area. We use openpyxl only for
    # reading; the final write still goes through SafeEditor/OOXML.
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=False, keep_vba=True)
    ws = wb[target_sheet]
    target_map: dict[str, int] = {}
    for row_idx in range(1, ws.max_row + 1):
        value = ws[f"{target_key_column.upper()}{row_idx}"].value
        if value is None:
            continue
        target_map[_norm_key(value)] = row_idx
    changes: list[BulkChange] = []
    missing: list[str] = []
    ops: list[EditOperation] = []
    for src in source_rows:
        key = _norm_key(src.get(source_key_field))
        if not key:
            missing.append("")
            continue
        row = target_map.get(key)
        if not row:
            missing.append(key)
            continue
        for target_col, source_field in mapping.items():
            raw = src.get(source_field)
            raw_text = "" if raw is None else str(raw)
            value, formula = _coerce(raw_text)
            cell = f"{target_col.upper()}{row}"
            snap = cell_info(p, target_sheet, cell)
            after = formula if formula is not None else value
            changes.append(BulkChange(target_sheet, cell, snap.formula if snap.formula is not None else snap.value, after, "formula" if formula else "value"))
            ops.append(EditOperation(cell, value=value, formula=formula))
    return {
        "target_sheet": target_sheet,
        "target_key_column": target_key_column.upper(),
        "source_key_field": source_key_field,
        "mapping": mapping,
        "matched_records": len(source_rows) - len(missing),
        "missing_keys": missing,
        "changes": [c.as_dict() for c in changes],
        "operations": ops,
    }


def _norm_key(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def keyed_execute(path: str | Path, target_sheet: str, target_key_column: str, source_rows: list[dict[str, Any]], source_key_field: str, mapping: dict[str, str], output: str | Path | None = None, *, create_backup: bool = True) -> dict[str, Any]:
    plan = keyed_preview(path, target_sheet, target_key_column, source_rows, source_key_field, mapping)
    if plan["missing_keys"]:
        return {k: v for k, v in plan.items() if k != "operations"} | {"executed": False, "reason": "missing_keys"}
    report = SafeEditor().edit_cells(path, target_sheet, plan["operations"], output, create_backup=create_backup)
    result = report.as_dict()
    result.update({k: v for k, v in plan.items() if k != "operations"})
    result["executed"] = True
    return result
